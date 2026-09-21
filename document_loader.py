from pathlib import Path

from pypdf import PdfReader
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def extract_pdf_text(file_path):
    """Extract text from a PDF file."""
    text_parts = []
    reader = PdfReader(file_path)

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text()

            if page_text:
                text_parts.append(
                    f"\n--- Page {page_number} ---\n{page_text}"
                )
        except Exception as e:
            print(f"Could not read page {page_number}: {e}")

    return "\n".join(text_parts)


def extract_docx_text(file_path):
    """Extract text from a DOCX file."""
    document = Document(file_path)
    text_parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            text_parts.append(text)

    # Also read tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text_parts.append(" | ".join(row_text))

    return "\n".join(text_parts)


def extract_txt_text(file_path):
    """Extract text from a TXT file."""
    with open(file_path, "r", encoding="utf-8") as file:
        return file.read()


def extract_text(file_path):
    """Detect file type and extract its text."""
    extension = Path(file_path).suffix.lower()

    if extension == ".pdf":
        return extract_pdf_text(file_path)

    if extension == ".docx":
        return extract_docx_text(file_path)

    if extension == ".txt":
        return extract_txt_text(file_path)

    raise ValueError(f"Unsupported file type: {extension}")


def load_documents(data_folder="data"):
    """Load all supported documents from the data folder."""
    documents = []
    data_path = Path(data_folder)

    if not data_path.exists():
        return documents

    for file_path in sorted(data_path.iterdir()):
        if not file_path.is_file():
            continue

        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        try:
            text = extract_text(file_path)

            if text.strip():
                documents.append(
                    {
                        "filename": file_path.name,
                        "file_path": str(file_path),
                        "text": text.strip(),
                    }
                )

        except Exception as e:
            print(f"Error reading {file_path.name}: {e}")

    return documents
